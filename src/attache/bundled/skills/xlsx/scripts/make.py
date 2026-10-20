# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Make an Excel workbook from a small JSON spec (SKILL.md describes it).

    make.py <spec> <out.xlsx>

<spec> is the spec itself, as JSON text starting with "{", or the name of a .json file. Both
files must be in the current folder, the chat's files folder when the script tool runs this:
absolute paths, drives, ":", ".." and hidden names are refused, and so is anything that
resolves outside the folder (through a link, say). The output must end in .xlsx. On success
the file's path relative to the Attaché home is printed, ready for `send`.

openpyxl is imported inside the functions that use it, so the checks work (and are tested)
without it.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path, PureWindowsPath
from typing import Any

SHEET = {"name", "rows", "header", "widths", "formats"}
DATE = re.compile(r"\d{4}-\d{2}-\d{2}(T\d{2}:\d{2}(:\d{2})?)?")  # becomes a real date
FIT = (8, 60)  # the narrowest and widest fitted column, in characters
LIVE = ("=", "+", "-", "@")  # text a spreadsheet may take for a formula when it is typed
# The functions a formula may use: arithmetic, aggregates and logic. Nothing that links,
# fetches, looks a cell up by name or reaches another program (HYPERLINK, WEBSERVICE,
# INDIRECT, ...), because a sheet a member asked for is opened by the owner or the group.
FUNCTIONS = {
    "SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "MEDIAN", "PRODUCT",
    "ROUND", "ROUNDUP", "ROUNDDOWN", "ABS",
    "IF", "IFERROR", "AND", "OR", "NOT",
    "SUMIF", "COUNTIF", "AVERAGEIF",
    "TODAY", "DATE",
}  # fmt: skip
CALL = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")  # a function name and its "("


def inside(name: str, suffix: str) -> Path:
    """`name` as a file in the current folder with the given suffix, or ValueError."""
    here = Path.cwd().resolve()
    windows = PureWindowsPath(name)  # splits on / and \ and sees drives, on every system
    hidden = any(part.startswith(".") for part in windows.parts)  # also ".."
    if not name or ":" in name or windows.anchor or hidden:
        raise ValueError(f"{name!r}: give a file name in this folder, such as budget{suffix}")
    path = here.joinpath(*windows.parts).resolve()
    if not path.is_relative_to(here):
        raise ValueError(f"{name!r} is outside this folder")
    if path.suffix.lower() != suffix:
        raise ValueError(f"{name!r}: the file must end in {suffix}")
    return path


def load(spec: str) -> dict[str, Any]:
    """The spec: JSON text, or the name of a .json file in the current folder."""
    if spec.lstrip().startswith("{"):
        text = spec
    else:
        text = inside(spec, ".json").read_text(encoding="utf-8")
    data = json.loads(text)
    if not isinstance(data, dict):
        raise ValueError("the spec must be a JSON object")
    return data


def shown(path: Path) -> str:
    """The path relative to the Attaché home (files/<chat>/budget.xlsx), as `send` takes it."""
    home = os.environ.get("ATTACHE_HOME")
    if home and path.is_relative_to(Path(home).resolve()):
        return path.relative_to(Path(home).resolve()).as_posix()
    base = next((parent.parent for parent in path.parents if parent.name == "files"), None)
    return (path.relative_to(base) if base else path).as_posix()


def check_keys(found: dict[str, Any], allowed: set[str], where: str) -> None:
    if unknown := sorted(set(found) - allowed):
        raise ValueError(f"{where}: unknown key {', '.join(unknown)}; known: {sorted(allowed)}")


def check_formula(formula: str) -> None:
    """ValueError naming the problem, for a formula that uses more than FUNCTIONS. A name
    inside quoted text counts too, which only ever refuses more."""
    if "|" in formula:
        raise ValueError(f"formula {formula}: | is not allowed (it runs another program)")
    if "[" in formula:
        raise ValueError(f"formula {formula}: [ is not allowed (it reads another workbook)")
    for name in CALL.findall(formula):
        if name.upper() not in FUNCTIONS:
            allowed = ", ".join(sorted(FUNCTIONS))
            raise ValueError(f"formula {formula}: {name} is not allowed; use only {allowed}")


def write(target: Any, value: Any) -> None:
    """Put one JSON value in a cell. Only {"formula": "=SUM(B2:B4)"} makes a formula, using
    FUNCTIONS alone: text is always text, even when it starts with "=", because it may come
    from someone's message (=HYPERLINK or =WEBSERVICE would run in Excel). ISO dates become
    dates; numbers, true/false and null stay as they are."""
    if isinstance(value, dict):
        formula = value.get("formula")
        if set(value) != {"formula"} or not isinstance(formula, str) or not formula.strip("= "):
            raise ValueError(f'cell {value!r}: a formula is {{"formula": "=SUM(B2:B4)"}}')
        check_formula(formula)
        target.value = formula if formula.startswith("=") else f"={formula}"
    elif isinstance(value, str) and DATE.fullmatch(value):
        moment = datetime.fromisoformat(value)
        target.value = moment if "T" in value else moment.date()
    elif value is None or isinstance(value, str | int | float):
        target.value = value
        if isinstance(value, str) and value.startswith(LIVE):
            target.data_type = "s"  # openpyxl took "=..." for a formula: keep it text
            target.quotePrefix = True  # Excel's leading ', so editing the cell keeps it text
    else:
        raise ValueError(f"cell {value!r}: use text, a number, true/false, null or a formula")


def by_column(table: Any, what: str) -> dict[int, Any]:
    """{"B": value} as {2: value}; ValueError for anything but column letters."""
    from openpyxl.utils import column_index_from_string

    if not isinstance(table, dict):
        raise ValueError(f'"{what}" maps column letters to values, like {{"A": 20}}')
    return {column_index_from_string(str(key).upper()): value for key, value in table.items()}


def fill(sheet: Any, spec: dict[str, Any]) -> None:
    """Write one sheet: its rows, then the bold frozen header, widths and number formats."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rows = spec.get("rows", [])
    if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
        raise ValueError(f"sheet {sheet.title!r}: rows is a list of rows, each a list of cells")
    for number, row in enumerate(rows, 1):
        for column, value in enumerate(row, 1):
            write(sheet.cell(row=number, column=column), value)
    header = bool(spec.get("header", True)) and bool(rows)
    if header:
        for column in range(1, len(rows[0]) + 1):
            sheet.cell(row=1, column=column).font = Font(bold=True)
        sheet.freeze_panes = "A2"  # the header stays in view when scrolling
    widths = by_column(spec.get("widths", {}), "widths")
    for column in range(1, max(map(len, rows), default=0) + 1):
        values = [row[column - 1] for row in rows if len(row) >= column]
        texts = [str(value) for value in values if not isinstance(value, dict | None)]
        longest = max(map(len, texts), default=0)  # a formula's text is not what it shows
        fitted = min(max(longest + 2, FIT[0]), FIT[1])
        width = widths.get(column, fitted)  # "widths" wins
        sheet.column_dimensions[get_column_letter(column)].width = float(width)
    for column, number_format in by_column(spec.get("formats", {}), "formats").items():
        for row in range(2 if header else 1, len(rows) + 1):
            sheet.cell(row=row, column=column).number_format = str(number_format)


def build(spec: dict[str, Any]) -> Any:
    """The openpyxl Workbook for `spec`."""
    import openpyxl

    check_keys(spec, {"sheets"}, "the spec")
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError('"sheets" is a list of at least one sheet')
    book = openpyxl.Workbook()
    book.remove(book.active)  # a new workbook starts with an empty sheet
    for number, sheet in enumerate(sheets, 1):
        if not isinstance(sheet, dict):
            raise ValueError(f"sheet {number}: a sheet is an object with rows")
        check_keys(sheet, SHEET, f"sheet {number}")
        fill(book.create_sheet(title=str(sheet.get("name") or f"Sheet{number}")), sheet)
    return book


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Make an Excel workbook from a JSON spec.")
    parser.add_argument("spec", help='the spec as JSON text ("{...}"), or a .json file here')
    parser.add_argument("out", help="the .xlsx file to write in the current folder")
    args = parser.parse_args(argv)
    try:
        out = inside(args.out, ".xlsx")
        book = build(load(args.spec))
        out.parent.mkdir(parents=True, exist_ok=True)
        book.save(str(out))
    except (OSError, ValueError, TypeError, KeyError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    print(shown(out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
